# import demistomock as demisto
# from CommonServerPython import *
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
import os

def insert_image_to_excel(excel_path, image_path, cell="A1", sheet_name=None):
    # Baca 4 byte pertama dari file Excel
    with open(excel_path, 'rb') as f:
        header = f.read(4)

    # Pastikan kita mendapatkan representasi hex dari header dengan aman
    if isinstance(header, bytes):
        header_hex = " ".join("{:02x}".format(b) for b in header)
    else:
        # Jika header berupa string, gunakan ord() untuk mendapatkan nilai integer tiap karakter
        header_hex = " ".join("{:02x}".format(ord(b)) for b in header)

    demisto.debug("Excel file header (hex): " + header_hex)

    # Buka workbook dengan openpyxl
    wb = load_workbook(excel_path)
    if sheet_name:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            return_error("Sheet '{}' tidak ditemukan dalam workbook.".format(sheet_name))
    else:
        ws = wb.active

    # Buat objek gambar dan tambahkan ke cell yang ditentukan
    img = XLImage(image_path)
    ws.add_image(img, cell)
    wb.save(excel_path)

def main():
    try:
        # Ambil argumen dari War Room
        excel_entry = demisto.args().get('excelFileEntryID')
        image_entry = demisto.args().get('imageFileEntryID')
        cell = demisto.args().get('cellPosition', 'A1')
        sheet = demisto.args().get('sheetName')

        if not excel_entry or not image_entry:
            return_error("Argumen excelFileEntryID dan imageFileEntryID harus disediakan.")

        # Dapatkan path fisik file dari Entry ID
        excel_info = demisto.getFilePath(excel_entry)
        image_info = demisto.getFilePath(image_entry)

        if not excel_info or not excel_info.get('path'):
            return_error("Gagal mendapatkan path file Excel.")
        if not image_info or not image_info.get('path'):
            return_error("Gagal mendapatkan path file gambar.")

        excel_path = excel_info['path']
        image_path = image_info['path']

        # Validasi ekstensi file Excel
        supported_exts = ['.xlsx', '.xlsm', '.xltx', '.xltm']
        excel_filename = excel_info.get('name', os.path.basename(excel_path))
        ext = os.path.splitext(excel_filename)[1].lower()
        if ext not in supported_exts:
            return_error("Format file Excel '{}' tidak didukung. Format yang didukung: {}".format(ext, ", ".join(supported_exts)))

        # Sisipkan gambar ke file Excel
        insert_image_to_excel(excel_path, image_path, cell, sheet)

        # Baca file Excel yang sudah diperbarui
        with open(excel_path, 'rb') as f:
            updated_data = f.read()

        demisto.results(fileResult(filename="UpdatedExcel.xlsx", data=updated_data))

    except Exception as e:
        return_error("Error occurred: " + str(e))

if __name__ in ('__main__', '__builtin__', 'builtins'):
    main()
